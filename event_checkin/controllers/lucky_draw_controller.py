import secrets
from io import BytesIO

from flask import Blueprint, jsonify, render_template, request, url_for
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from PIL import Image, UnidentifiedImageError
from sqlalchemy.exc import IntegrityError

from event_checkin.controllers.admin_controller import (
    _ascii_filename_component,
    _xlsx_response,
    admin_required,
)
from event_checkin.models import db
from event_checkin.models.checkin_record import CheckinRecord
from event_checkin.models.event import Event
from event_checkin.models.lucky_draw import LuckyConfig, LuckyPrize, LuckyWinner
from event_checkin.models.user import User
from event_checkin.utils.lucky_draw import UPLOAD_SUBDIR, background_path, remove_files, upload_dir
from event_checkin.utils.timezone import format_local, now_vn_naive


lucky_bp = Blueprint("lucky", __name__)

DEFAULT_BACKGROUND = "lucky/background.jpg"
BACKGROUND_MAX_BYTES = 5 * 1024 * 1024
BACKGROUND_MAX_SIDE = 8000
BACKGROUND_FORMATS = {"JPEG": "jpg", "PNG": "png", "WEBP": "webp"}
MAX_PRIZE_NAME_LENGTH = 255
MAX_PRIZE_QUANTITY = 1000
DECOY_LIMIT = 100  # số mã "mồi" gửi cho màn hình để chạy hiệu ứng
TIME_FORMAT = "%H:%M:%S %d/%m/%Y"


class DrawError(Exception):
    def __init__(self, message, status=400):
        super().__init__(message)
        self.message = message
        self.status = status


@lucky_bp.errorhandler(DrawError)
def _handle_draw_error(error):
    db.session.rollback()
    return jsonify({"success": False, "message": error.message}), error.status


# ---------- Dữ liệu dùng chung ----------

def _ensure_config(event_id):
    """Đảm bảo sự kiện có dòng cấu hình (commit riêng, để dòng này tồn tại trước khi khoá)."""
    if db.session.get(LuckyConfig, event_id):
        db.session.commit()  # kết thúc transaction đọc
        return
    db.session.add(LuckyConfig(event_id=event_id, cho_phep_trung_nhieu_giai=False))
    try:
        db.session.commit()
    except IntegrityError:  # tab khác vừa tạo -- không sao
        db.session.rollback()


def _lock_config(event_id):
    """Khoá hàng cấu hình của sự kiện đến hết transaction (SELECT ... FOR UPDATE). Mọi thao tác
    quay/lưu/huỷ đều đi qua đây nên chạy tuần tự theo từng sự kiện.

    MySQL mặc định REPEATABLE READ: ảnh chụp dữ liệu được lấy ở lần đọc đầu tiên của transaction.
    admin_required đã đọc bảng admins từ trước, nên nếu cứ khoá tiếp trong transaction đó thì sau
    khi chờ được khoá, tab thứ hai vẫn thấy dữ liệu CŨ (không thấy kết quả chờ của tab thứ nhất).
    Vì vậy: kết thúc transaction đọc, mở transaction mới ở mức READ COMMITTED rồi mới khoá --
    mọi lần đọc sau khoá đều thấy dữ liệu mới nhất."""
    db.session.commit()
    _ensure_config(event_id)  # có thể commit; phải làm trước khi đặt mức cô lập cho transaction khoá
    if db.engine.dialect.name in ("mysql", "mariadb"):
        db.session.connection(execution_options={"isolation_level": "READ COMMITTED"})
    return LuckyConfig.query.filter_by(event_id=event_id).with_for_update().populate_existing().one()


def _valid_code(code):
    return bool(code) and len(code) <= 50 and not any(ch.isspace() for ch in code)


def _participants(event_id):
    """{patron_code: họ tên} của mọi người có ít nhất 1 dòng check-in, gộp trùng theo mã.
    KHÔNG lọc theo cert_status: lỗi gửi email không liên quan đến việc người đó có mặt.
    Tên: cert_display_name -> patron_name -> tên đăng ký (User) -> trống; ưu tiên dòng không 'duplicate'."""
    records = (
        CheckinRecord.query.filter_by(event_id=event_id)
        .order_by(CheckinRecord.checkin_at.asc(), CheckinRecord.id.asc())
        .all()
    )
    by_code = {}
    for item in records:
        code = (item.patron_code or "").strip()
        if _valid_code(code):
            by_code.setdefault(code, []).append(item)
    users = {
        user.ma_cbsv: user
        for user in User.query.filter(User.ma_cbsv.in_(list(by_code))).all()
    } if by_code else {}
    result = {}
    for code, items in by_code.items():
        items = sorted(items, key=lambda record: record.cert_status == "duplicate")
        name = (
            next((i.cert_display_name.strip() for i in items if i.cert_display_name and i.cert_display_name.strip()), "")
            or next((i.patron_name.strip() for i in items if i.patron_name and i.patron_name.strip()), "")
            or (users[code].ho_ten if code in users else "")
        )
        result[code] = name
    return result


def _active_winners(event_id):
    return LuckyWinner.query.filter_by(event_id=event_id, hieu_luc=1).all()


def _eligible_codes(participants, allow_multi, active, prize_id):
    """Mã còn được quay cho giải prize_id: bỏ người đã trúng (mọi giải, hoặc chỉ giải này nếu
    cho phép trúng nhiều giải). Tính cả kết quả đang chờ Lưu."""
    if allow_multi:
        taken = {w.patron_code for w in active if w.prize_id == prize_id}
    else:
        taken = {w.patron_code for w in active}
    return sorted(code for code in participants if code not in taken)


def _background_url(config):
    path = background_path(config.hinh_nen) if config else None
    if path and path.exists():
        return url_for("static", filename=config.hinh_nen)
    return url_for("static", filename=DEFAULT_BACKGROUND)


def _winner_payload(winner, with_name):
    payload = {
        "id": winner.id,
        "prize_id": winner.prize_id,
        "ten_giai": winner.prize.ten_giai if winner.prize else "",
        "patron_code": winner.patron_code,
        "trang_thai": winner.trang_thai,
        "trung_luc": format_local(winner.trung_luc, TIME_FORMAT),
    }
    if with_name:  # màn hình chiếu chỉ cần mã, không gửi họ tên
        payload["ho_ten"] = winner.ho_ten or ""
        payload["huy_luc"] = format_local(winner.huy_luc, TIME_FORMAT)
    return payload


def _state(event, with_name):
    config = db.session.get(LuckyConfig, event.id)
    allow_multi = bool(config and config.cho_phep_trung_nhieu_giai)
    participants = _participants(event.id)
    prizes = LuckyPrize.query.filter_by(event_id=event.id).order_by(LuckyPrize.thu_tu, LuckyPrize.id).all()
    winners = (
        LuckyWinner.query.filter_by(event_id=event.id)
        .order_by(LuckyWinner.trung_luc.asc(), LuckyWinner.id.asc())
        .all()
    )
    active = [w for w in winners if w.hieu_luc == 1]
    prize_rows = []
    for prize in prizes:
        used = sum(1 for w in active if w.prize_id == prize.id)
        prize_rows.append({
            "id": prize.id,
            "ten_giai": prize.ten_giai,
            "so_luong": prize.so_luong,
            "thu_tu": prize.thu_tu,
            "da_quay": used,
            "con_lai": max(prize.so_luong - used, 0),
            "nguoi_con_lai": len(_eligible_codes(participants, allow_multi, active, prize.id)),
        })
    return {
        "success": True,
        "event": {"id": event.id, "ten_su_kien": event.ten_su_kien},
        "tong_nguoi": len(participants),
        "cho_phep_trung_nhieu_giai": allow_multi,
        "hinh_nen_url": _background_url(config),
        "hinh_nen_tuy_chinh": bool(config and config.hinh_nen),
        "prizes": prize_rows,
        "pending": [_winner_payload(w, with_name) for w in winners if w.trang_thai == "pending"],
        "winners": [_winner_payload(w, with_name) for w in winners if w.trang_thai == "confirmed"],
        "cancelled": [_winner_payload(w, with_name) for w in winners if w.trang_thai == "cancelled"] if with_name else [],
    }


# ---------- Trang ----------

@lucky_bp.get("/admin/events/<int:event_id>/lucky-draw")
@admin_required
def lucky_page(event_id):
    event = Event.query.get_or_404(event_id)
    return render_template(
        "admin/lucky_draw.html",
        event=event,
        back_url=url_for("admin.event_detail", event_id=event.id),
        screen_url=url_for("lucky.lucky_screen", event_id=event.id),
        default_background_url=url_for("static", filename=DEFAULT_BACKGROUND),
    )


@lucky_bp.get("/admin/events/<int:event_id>/lucky-draw/screen")
@admin_required
def lucky_screen(event_id):
    event = Event.query.get_or_404(event_id)
    return render_template(
        "admin/lucky_screen.html",
        event=event,
        back_url=url_for("lucky.lucky_page", event_id=event.id),
    )


# ---------- Trạng thái / cấu hình ----------

@lucky_bp.get("/api/admin/events/<int:event_id>/lucky-draw")
@admin_required
def lucky_state(event_id):
    event = Event.query.get_or_404(event_id)
    with_name = request.args.get("screen") != "1"
    return jsonify(_state(event, with_name))


@lucky_bp.put("/api/admin/events/<int:event_id>/lucky-draw/settings")
@admin_required
def lucky_save_settings(event_id):
    Event.query.get_or_404(event_id)
    payload = request.get_json(silent=True) or {}
    if not isinstance(payload.get("cho_phep_trung_nhieu_giai"), bool):
        raise DrawError("Giá trị cài đặt không hợp lệ.")
    config = _lock_config(event_id)
    config.cho_phep_trung_nhieu_giai = payload["cho_phep_trung_nhieu_giai"]
    db.session.commit()
    return jsonify({"success": True, "message": "Đã lưu cài đặt."})


# ---------- Giải thưởng ----------

def _parse_int(value, field, minimum, maximum):
    if isinstance(value, bool):
        raise DrawError(f"{field} không hợp lệ.")
    try:
        number = int(value)
    except (TypeError, ValueError):
        raise DrawError(f"{field} không hợp lệ.")
    if number < minimum or number > maximum:
        raise DrawError(f"{field} phải từ {minimum} đến {maximum}.")
    return number


def _prize_name(value):
    name = " ".join(str(value or "").split())
    if not name:
        raise DrawError("Vui lòng nhập tên giải.")
    if len(name) > MAX_PRIZE_NAME_LENGTH:
        raise DrawError(f"Tên giải tối đa {MAX_PRIZE_NAME_LENGTH} ký tự.")
    return name


def _get_prize(event_id, prize_id):
    prize = db.session.get(LuckyPrize, prize_id)
    if not prize or prize.event_id != event_id:
        raise DrawError("Không tìm thấy giải thưởng.", 404)
    return prize


@lucky_bp.post("/api/admin/events/<int:event_id>/lucky-draw/prizes")
@admin_required
def lucky_create_prize(event_id):
    Event.query.get_or_404(event_id)
    payload = request.get_json(silent=True) or {}
    name = _prize_name(payload.get("ten_giai"))
    quantity = _parse_int(payload.get("so_luong", 1), "Số lượng", 1, MAX_PRIZE_QUANTITY)
    _lock_config(event_id)
    if payload.get("thu_tu") in (None, ""):
        last = db.session.query(db.func.max(LuckyPrize.thu_tu)).filter_by(event_id=event_id).scalar()
        order = (last or 0) + 1
    else:
        order = _parse_int(payload.get("thu_tu"), "Thứ tự", 0, 100000)
    db.session.add(LuckyPrize(event_id=event_id, ten_giai=name, so_luong=quantity, thu_tu=order))
    db.session.commit()
    return jsonify({"success": True, "message": "Đã thêm giải."})


@lucky_bp.put("/api/admin/events/<int:event_id>/lucky-draw/prizes/<int:prize_id>")
@admin_required
def lucky_update_prize(event_id, prize_id):
    Event.query.get_or_404(event_id)
    payload = request.get_json(silent=True) or {}
    _lock_config(event_id)
    prize = _get_prize(event_id, prize_id)
    if "ten_giai" in payload:
        prize.ten_giai = _prize_name(payload["ten_giai"])
    if "so_luong" in payload:
        quantity = _parse_int(payload["so_luong"], "Số lượng", 1, MAX_PRIZE_QUANTITY)
        used = LuckyWinner.query.filter_by(prize_id=prize.id, hieu_luc=1).count()
        if quantity < used:
            raise DrawError(f"Giải này đã có {used} người trúng, không thể giảm số lượng xuống {quantity}.")
        prize.so_luong = quantity
    if "thu_tu" in payload:
        prize.thu_tu = _parse_int(payload["thu_tu"], "Thứ tự", 0, 100000)
    db.session.commit()
    return jsonify({"success": True, "message": "Đã cập nhật giải."})


@lucky_bp.delete("/api/admin/events/<int:event_id>/lucky-draw/prizes/<int:prize_id>")
@admin_required
def lucky_delete_prize(event_id, prize_id):
    Event.query.get_or_404(event_id)
    _lock_config(event_id)
    prize = _get_prize(event_id, prize_id)
    if LuckyWinner.query.filter_by(prize_id=prize.id).count():
        raise DrawError("Giải này đã có lịch sử quay (kể cả bản huỷ) nên không thể xoá. Hãy đổi tên hoặc sắp xếp lại thay vì xoá.", 409)
    db.session.delete(prize)
    db.session.commit()
    return jsonify({"success": True, "message": "Đã xoá giải."})


# ---------- Quay / Lưu / Quay lại / Huỷ ----------

@lucky_bp.post("/api/admin/events/<int:event_id>/lucky-draw/draw")
@admin_required
def lucky_draw(event_id):
    Event.query.get_or_404(event_id)
    payload = request.get_json(silent=True) or {}
    prize_id = _parse_int(payload.get("prize_id"), "Giải", 1, 2**31 - 1)
    count = _parse_int(payload.get("so_nguoi", 1), "Số người mỗi lượt", 1, MAX_PRIZE_QUANTITY)

    config = _lock_config(event_id)  # từ đây các lượt quay của sự kiện này chạy lần lượt
    if LuckyWinner.query.filter_by(event_id=event_id, trang_thai="pending").count():
        raise DrawError("Đang có kết quả chờ Lưu hoặc Quay lại (có thể ở tab khác). Hãy xử lý kết quả đó trước.", 409)
    prize = _get_prize(event_id, prize_id)

    active = _active_winners(event_id)
    remaining = prize.so_luong - sum(1 for w in active if w.prize_id == prize.id)
    if remaining <= 0:
        raise DrawError(f"{prize.ten_giai} đã quay đủ {prize.so_luong} người.", 409)
    if count > remaining:
        raise DrawError(f"{prize.ten_giai} chỉ còn {remaining} suất, không thể quay {count} người.", 409)

    participants = _participants(event_id)
    if not participants:
        raise DrawError("Sự kiện này chưa có ai check-in nên chưa có danh sách để quay.", 409)
    pool = _eligible_codes(participants, config.cho_phep_trung_nhieu_giai, active, prize.id)
    if len(pool) < count:
        raise DrawError(f"Chỉ còn {len(pool)} người trong danh sách quay, không đủ để quay {count} người.", 409)

    chooser = secrets.SystemRandom()
    picks = chooser.sample(pool, count)
    now = now_vn_naive()
    winners = [
        LuckyWinner(
            event_id=event_id, prize_id=prize.id, patron_code=code,
            ho_ten=participants[code] or None, trang_thai="pending", trung_luc=now, hieu_luc=1,
        )
        for code in picks
    ]
    db.session.add_all(winners)
    try:
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        raise DrawError("Kết quả vừa bị trùng với một lượt quay khác, vui lòng thử lại.", 409)

    decoys = chooser.sample(sorted(participants), min(DECOY_LIMIT, len(participants)))
    return jsonify({
        "success": True,
        "winners": [{"id": w.id, "patron_code": w.patron_code} for w in winners],
        "decoys": decoys,
        "prize": {"id": prize.id, "ten_giai": prize.ten_giai},
    })


@lucky_bp.post("/api/admin/events/<int:event_id>/lucky-draw/confirm")
@admin_required
def lucky_confirm(event_id):
    """Lưu: chốt cả lượt đang chờ."""
    Event.query.get_or_404(event_id)
    _lock_config(event_id)
    pending = LuckyWinner.query.filter_by(event_id=event_id, trang_thai="pending").all()
    if not pending:
        raise DrawError("Không có kết quả nào đang chờ Lưu.", 409)
    for winner in pending:
        winner.trang_thai = "confirmed"
    db.session.commit()
    return jsonify({"success": True, "message": f"Đã lưu {len(pending)} người trúng thưởng."})


@lucky_bp.post("/api/admin/events/<int:event_id>/lucky-draw/reject")
@admin_required
def lucky_reject(event_id):
    """Quay lại: bỏ cả lượt đang chờ (chưa Lưu nên xoá hẳn, những người này quay lại danh sách quay)."""
    Event.query.get_or_404(event_id)
    _lock_config(event_id)
    count = LuckyWinner.query.filter_by(event_id=event_id, trang_thai="pending").delete()
    if not count:
        raise DrawError("Không có kết quả nào đang chờ để bỏ.", 409)
    db.session.commit()
    return jsonify({"success": True, "message": "Đã bỏ kết quả, có thể quay lại."})


@lucky_bp.post("/api/admin/events/<int:event_id>/lucky-draw/winners/<int:winner_id>/cancel")
@admin_required
def lucky_cancel(event_id, winner_id):
    """Huỷ người vắng mặt: giữ bản ghi để đối chiếu, người này quay lại danh sách quay."""
    Event.query.get_or_404(event_id)
    _lock_config(event_id)
    winner = db.session.get(LuckyWinner, winner_id)
    if not winner or winner.event_id != event_id:
        raise DrawError("Không tìm thấy kết quả.", 404)
    if winner.trang_thai != "confirmed":
        raise DrawError("Chỉ huỷ được kết quả đã lưu. Kết quả đang chờ hãy bấm Quay lại.", 409)
    winner.trang_thai = "cancelled"
    winner.huy_luc = now_vn_naive()
    winner.hieu_luc = None
    db.session.commit()
    return jsonify({"success": True, "message": f"Đã huỷ kết quả của {winner.patron_code}. Người này được đưa lại vào danh sách quay."})


# ---------- Ảnh nền ----------

@lucky_bp.post("/api/admin/events/<int:event_id>/lucky-draw/background")
@admin_required
def lucky_upload_background(event_id):
    Event.query.get_or_404(event_id)
    file = request.files.get("hinh_nen")
    if not file or not file.filename:
        raise DrawError("Vui lòng chọn hình nền.")
    data = file.stream.read(BACKGROUND_MAX_BYTES + 1)
    if len(data) > BACKGROUND_MAX_BYTES:
        raise DrawError("Hình nền tối đa 5MB.")
    try:
        with Image.open(BytesIO(data)) as image:
            detected = image.format
            width, height = image.size
            image.verify()  # kiểm tra đúng là file ảnh nguyên vẹn, không tin đuôi file
    except (UnidentifiedImageError, OSError, SyntaxError, Image.DecompressionBombError):
        raise DrawError("File không phải ảnh hợp lệ. Chỉ nhận JPG, PNG hoặc WEBP.")
    extension = BACKGROUND_FORMATS.get(detected)
    if not extension:
        raise DrawError("Chỉ nhận ảnh JPG, PNG hoặc WEBP.")
    if width > BACKGROUND_MAX_SIDE or height > BACKGROUND_MAX_SIDE:
        raise DrawError(f"Kích thước ảnh tối đa {BACKGROUND_MAX_SIDE}px mỗi cạnh.")

    folder = upload_dir()
    folder.mkdir(parents=True, exist_ok=True)
    filename = f"bg_{event_id}_{secrets.token_hex(12)}.{extension}"  # tên ngẫu nhiên, không dùng tên người dùng
    (folder / filename).write_bytes(data)

    config = _lock_config(event_id)
    old_path = background_path(config.hinh_nen)
    config.hinh_nen = f"{UPLOAD_SUBDIR}/{filename}"
    db.session.commit()
    remove_files([old_path] if old_path else [])
    return jsonify({"success": True, "message": "Đã đổi hình nền.", "hinh_nen_url": _background_url(config)})


@lucky_bp.delete("/api/admin/events/<int:event_id>/lucky-draw/background")
@admin_required
def lucky_reset_background(event_id):
    Event.query.get_or_404(event_id)
    config = _lock_config(event_id)
    old_path = background_path(config.hinh_nen)
    config.hinh_nen = None
    db.session.commit()
    remove_files([old_path] if old_path else [])
    return jsonify({"success": True, "message": "Đã khôi phục hình nền mặc định.", "hinh_nen_url": _background_url(config)})


# ---------- Xuất Excel ----------

@lucky_bp.get("/api/admin/events/<int:event_id>/lucky-draw/export")
@admin_required
def lucky_export(event_id):
    event = Event.query.get_or_404(event_id)
    winners = LuckyWinner.query.filter(LuckyWinner.event_id == event_id, LuckyWinner.trang_thai != "pending").all()
    confirmed = sorted(
        (w for w in winners if w.trang_thai == "confirmed"),
        key=lambda w: (w.prize.thu_tu, w.prize_id, w.trung_luc, w.id),
    )
    if not confirmed:
        return jsonify({"success": False, "message": "Chưa có kết quả trúng thưởng để xuất."}), 400
    cancelled = sorted((w for w in winners if w.trang_thai == "cancelled"), key=lambda w: (w.huy_luc or w.trung_luc, w.id))

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Trung thuong"
    sheet.append(["STT", "Giải", "Mã CB/SV", "Họ tên", "Giờ trúng"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for position, w in enumerate(confirmed, start=1):
        sheet.append([position, w.prize.ten_giai, w.patron_code, w.ho_ten or "", format_local(w.trung_luc, TIME_FORMAT)])
    for index, width in enumerate([6, 24, 16, 28, 22], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    if cancelled:
        sheet2 = workbook.create_sheet("Da huy (vang mat)")
        sheet2.append(["STT", "Giải", "Mã CB/SV", "Họ tên", "Giờ trúng", "Giờ huỷ"])
        for cell in sheet2[1]:
            cell.font = Font(bold=True)
        for position, w in enumerate(cancelled, start=1):
            sheet2.append([
                position, w.prize.ten_giai, w.patron_code, w.ho_ten or "",
                format_local(w.trung_luc, TIME_FORMAT), format_local(w.huy_luc, TIME_FORMAT),
            ])
        for index, width in enumerate([6, 24, 16, 28, 22, 22], start=1):
            sheet2.column_dimensions[get_column_letter(index)].width = width

    today = now_vn_naive().strftime("%d-%m-%Y")
    return _xlsx_response(
        workbook,
        f"{event.ten_su_kien}_quayso_{today}.xlsx",
        f"{_ascii_filename_component(event.ten_su_kien)}_quayso_{today}.xlsx",
    )
